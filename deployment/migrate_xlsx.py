#!/usr/bin/env python3
"""Prepare an XLSX snapshot locally and import it once into an empty DailyTask DB."""
import argparse
import base64
import datetime as dt
import hashlib
import json
import math
import os
import re
import secrets
import shutil
import sqlite3
from pathlib import Path

SHEETS = ('Users','Roles','Tasks','ContentPlan','Shoots','Lives','Channels','ChannelStats','Meetings','Shifts','Policy','Inputs','Payroll','AutoInputs')
KEYS = {
    'Users': ('email',), 'Roles': ('role',), 'Tasks': ('id',), 'ContentPlan': ('id',),
    'Shoots': ('id',), 'Lives': ('id',), 'Channels': ('email','slot'),
    'ChannelStats': ('month','email','slot'), 'Meetings': ('id',), 'Shifts': ('week','email'),
    'Policy': ('email','code'), 'Inputs': ('month','email'), 'Payroll': ('month','email'),
    'AutoInputs': ('month','email'),
}

def val(x):
    if x is None: return None
    if isinstance(x, dt.datetime):
        return x.date().isoformat() if x.time() == dt.time() else x.isoformat(sep=' ', timespec='seconds')
    if isinstance(x, (dt.date, dt.time)): return x.isoformat()
    if isinstance(x, float) and not math.isfinite(x): raise ValueError('non-finite numeric value')
    return x if isinstance(x, (bool,int,float)) else str(x)

def text(x, default=''):
    v=val(x)
    return default if v is None else str(v).strip()

def month(x):
    s=text(x)
    return s[:7] if s else ''

def flag(x, default=False):
    if x is None or x == '': return int(default)
    if isinstance(x, bool): return int(x)
    if isinstance(x, (int,float)): return int(x != 0)
    s=str(x).strip().lower()
    if s in ('true','1','yes','có','active','đang hoạt động'): return 1
    if s in ('false','0','no','không','inactive','đã nghỉ'): return 0
    raise ValueError(f'invalid boolean value: {s}')

def split_list(x):
    return [v.strip() for v in text(x).split(',') if v.strip()]

def normalize_json(x, default):
    s=text(x)
    if not s: return json.dumps(default, ensure_ascii=False)
    try: return json.dumps(json.loads(s), ensure_ascii=False, separators=(',',':'))
    except (json.JSONDecodeError, TypeError): return s

def read_book(path):
    import openpyxl
    book=openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing=set(SHEETS)-set(book.sheetnames)
    if missing: raise ValueError('missing expected sheets: '+', '.join(sorted(missing)))
    data={}
    for name in SHEETS:
        rows=book[name].iter_rows(values_only=True)
        header=next(rows, None)
        if not header or not any(v not in (None,'') for v in header):
            data[name]=[]; continue
        cols=[text(v) for v in header]
        if len(cols)!=len(set(cols)): raise ValueError(f'{name}: duplicate header')
        data[name]=[{cols[i]:r[i] for i in range(min(len(cols),len(r)))} for r in rows if any(v not in (None,'') for v in r)]
    return data

def prepare(source, destination):
    d=read_book(source)
    for name,keys in KEYS.items():
        seen=set()
        for n,r in enumerate(d[name],2):
            key=tuple(text(r.get(k)).lower() for k in keys)
            if any(not x for x in key): raise ValueError(f'{name} row {n}: blank primary key')
            if key in seen: raise ValueError(f'{name} row {n}: duplicate key')
            seen.add(key)
    roles={text(r.get('role')).lower() for r in d['Roles']}
    users={text(r.get('email')).lower() for r in d['Users']}
    if not users or any('@' not in e for e in users): raise ValueError('Users: missing or invalid email')
    usernames=[text(r.get('username')).lower() for r in d['Users'] if text(r.get('username'))]
    if len(usernames)!=len(set(usernames)): raise ValueError('Users: duplicate username')
    for r in d['Users']:
        if text(r.get('role')).lower() not in roles: raise ValueError('Users: role is missing from Roles')
        salt=text(r.get('salt')); old=text(r.get('password_hash')); temp=text(r.get('temp_password'))
        if old and salt:
            try: base64.b64decode(old, validate=True)
            except Exception as e: raise ValueError('Users: invalid legacy password hash') from e
            r['__credential']='legacy-sha256$'+salt+'$'+old
        elif temp:
            salt=secrets.token_urlsafe(18)
            digest=base64.b64encode(hashlib.sha256((salt+temp).encode()).digest()).decode()
            r['__credential']='legacy-sha256$'+salt+'$'+digest
        else: raise ValueError('Users: account has no migratable credential')
        for key in ('temp_password','password_hash','salt'): r.pop(key,None)
        for key,value in list(r.items()): r[key]=val(value)
        r['__email']=text(r.get('email')).lower()
        r['__role']=text(r.get('role')).lower()
        r['__active']=flag(r.get('active'),True)
        r['__username']=text(r.get('username')).lower() or None
    for name,field in [('Tasks','assignee'),('ContentPlan','assignee'),('Shoots','assignee'),('Lives','host'),('Channels','email'),('Shifts','email'),('Inputs','email'),('Payroll','email')]:
        for n,r in enumerate(d[name],2):
            e=text(r.get(field)).lower()
            if e and e not in users: raise ValueError(f'{name} row {n}: user reference missing from Users')
    for name,rows in d.items():
        if name=='Users': continue
        for r in rows:
            for k,v in list(r.items()):
                if name=='Policy' and k=='tiers' and isinstance(v,dt.timedelta):
                    if text(r.get('code'))!='live_bonus' or text(r.get('input_key'))!='live_total':
                        raise ValueError('Policy: unexpected Excel time conversion in tiers')
                    # Excel auto-coerced the legacy "30:500000" threshold/reward to a duration.
                    r[k]='30:500000'
                else: r[k]=val(v)
    # Normalize month/date-backed Excel cells before serialization.
    for name in ('Inputs','Payroll','Policy','AutoInputs','ChannelStats'):
        for r in d[name]:
            if 'month' in r: r['month']=month(r['month'])
    payload={'source':str(Path(source).resolve()),'counts':{k:len(v) for k,v in d.items()},'data':d}
    out=Path(destination); out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(payload,ensure_ascii=False,separators=(',',':')))
    out.chmod(0o600)
    print('Prepared rows:',json.dumps(payload['counts'],ensure_ascii=False,sort_keys=True))
    print('Credential values are hashed; plaintext temporary passwords were excluded.')

def apply(payload_path, db_path):
    payload=json.loads(Path(payload_path).read_text())
    d=payload['data']; counts=payload['counts']; db=Path(db_path)
    con=sqlite3.connect(db_path, timeout=30)
    con.execute('PRAGMA foreign_keys=ON')
    try:
        empty={'tasks':'Tasks','content_plan':'ContentPlan','schedules':'Shoots+Lives','meetings':'Meetings','shifts':'Shifts','policies':'Policy','inputs':'Inputs','payroll':'Payroll','channels':'Channels','channel_stats':'ChannelStats'}
        for table,source in empty.items():
            n=con.execute(f'SELECT count(*) FROM {table}').fetchone()[0]
            if n: raise ValueError(f'{table} already contains {n} rows; refusing to merge/overwrite')
        existing={x[0].lower() for x in con.execute('SELECT email FROM users')}
        source_emails={r['__email'] for r in d['Users']}
        if existing != {'admin@example.com'} or existing & source_emails:
            raise ValueError('production users differ from the expected sole bootstrap admin; refusing to merge')
        known_roles={x[0] for x in con.execute('SELECT name FROM roles')}
        if not known_roles.issubset({text(r.get('role')).lower() for r in d['Roles']}): raise ValueError('production roles differ from the expected bootstrap roles')
        stamp=dt.datetime.now().strftime('%Y%m%d-%H%M%S')
        backup=db.parent/f'pre-migration-{stamp}.sqlite3'
        if backup.exists(): raise ValueError('backup path already exists')
        dest=sqlite3.connect(backup)
        con.backup(dest); dest.close(); backup.chmod(0o600)
        con.execute('BEGIN IMMEDIATE')
        for r in d['Roles']:
            caps=[x.strip() for x in text(r.get('perms')).split(',') if x.strip()]
            con.execute('INSERT INTO roles(name,label,caps,locked) VALUES(?,?,?,?) ON CONFLICT(name) DO UPDATE SET label=excluded.label,caps=excluded.caps,locked=excluded.locked',
                (text(r.get('role')).lower(),text(r.get('label')),json.dumps(caps,ensure_ascii=False),flag(r.get('locked'))))
        for r in d['Users']:
            email=r['__email']; username=r['__username']; role=r['__role']; name=text(r.get('name'))
            initials=text(r.get('initials'))
            if not email or not name: raise ValueError('Users: missing email/name')
            con.execute('INSERT INTO users(email,username,name,initials,position,color,role,active,start_date,password_hash) VALUES(?,?,?,?,?,?,?,?,?,?)',
                (email,username,name,initials,text(r.get('position')),text(r.get('color')) or '#7657E8',role,r['__active'],text(r.get('start_date')),r['__credential']))
        for r in d['Tasks']:
            con.execute('INSERT INTO tasks(id,title,assignee,due,priority,status,created_at,updated_at,kpi_key,qty,done_at,schedule_id,due_date) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (text(r.get('id')),text(r.get('title')),text(r.get('assignee')).lower(),text(r.get('due')),text(r.get('priority')) or 'Vừa',text(r.get('status')) or 'todo',text(r.get('created_at')) or dt.datetime.now().isoformat(sep=' ',timespec='seconds'),text(r.get('updated_at')),text(r.get('kpi_key')),val(r.get('qty')) or 1,text(r.get('done_at')),text(r.get('schedule_id')),text(r.get('due_date'))))
        for r in d['ContentPlan']:
            con.execute('INSERT INTO content_plan(id,channel,month,pillar,content_key,demo_date,post_date,status,message,assignee) VALUES(?,?,?,?,?,?,?,?,?,?)',
                (text(r.get('id')),text(r.get('channel')),month(r.get('month')),text(r.get('pillar')),text(r.get('key')),text(r.get('demo_date')),text(r.get('post_date')),text(r.get('status')) or 'Chưa thực hiện',text(r.get('message')),text(r.get('assignee')).lower()))
        for sheet,kind,title_key,lead_key in [('Shoots','shoot','topic','assignee'),('Lives','live','title','host')]:
            for r in d[sheet]:
                con.execute('INSERT INTO schedules(id,kind,title,date,time,location,lead,attendees,status,kpi_key,qty,done,brief,created_by,event_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                    (text(r.get('id')),kind,text(r.get(title_key)),text(r.get('date')),text(r.get('time')),text(r.get('location')),text(r.get(lead_key)).lower(),'[]',text(r.get('status')) or 'Đã lên lịch',text(r.get('kpi_key')),val(r.get('qty')) or 1,flag(r.get('done')),text(r.get('brief')),'admin@example.com',text(r.get('event_id'))))
        for r in d['Meetings']:
            attendees=split_list(r.get('attendee_emails')) or split_list(r.get('attendees'))
            duration=re.search(r'\d+',text(r.get('duration')))
            con.execute('INSERT INTO meetings(id,title,date,time,attendees,duration,repeat,note,created_by,event_id) VALUES(?,?,?,?,?,?,?,?,?,?)',
                (text(r.get('id')),text(r.get('title')),text(r.get('date')),text(r.get('time')),json.dumps(attendees,ensure_ascii=False),int(duration.group()) if duration else 60,text(r.get('repeat')),text(r.get('note')),'admin@example.com',text(r.get('event_id'))))
        for r in d['Shifts']:
            con.execute('INSERT INTO shifts(week,email,mon,tue,wed,thu,fri,sat,sun) VALUES(?,?,?,?,?,?,?,?,?)',
                (text(r.get('week'))[:10],text(r.get('email')).lower(),*(text(r.get(k)) for k in ('T2','T3','T4','T5','T6','T7','CN'))))
        for r in d['Policy']:
            con.execute('INSERT INTO policies(email,code,label,type,input_key,rate,tiers,minimum,note,active) VALUES(?,?,?,?,?,?,?,?,?,?)',
                (text(r.get('email')).lower(),text(r.get('code')),text(r.get('label')),text(r.get('type')),text(r.get('input_key')),val(r.get('rate')) or 0,text(r.get('tiers')),val(r.get('min')) or 0,text(r.get('note')),flag(r.get('active'),True)))
        for r in d['Inputs']:
            values={k:val(v) for k,v in r.items() if k not in ('month','email') and v not in (None,'')}
            con.execute('INSERT INTO inputs(month,email,values_json) VALUES(?,?,?)',(month(r.get('month')),text(r.get('email')).lower(),json.dumps(values,ensure_ascii=False,separators=(',',':'))))
        for r in d['Payroll']:
            con.execute('INSERT INTO payroll(month,email,base,fees,bonus,penalty,total,kpi_ok,kpi_met,kpi_total,kpi_rate,kpi_items,breakdown,computed_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (month(r.get('month')),text(r.get('email')).lower(),val(r.get('base')) or 0,val(r.get('fees')) or 0,val(r.get('bonus')) or 0,val(r.get('penalty')) or 0,val(r.get('total')) or 0,flag(r.get('kpi_ok'),True),val(r.get('kpi_met')) or 0,val(r.get('kpi_total')) or 0,val(r.get('kpi_rate')) if r.get('kpi_rate') not in (None,'') else 100,normalize_json(r.get('kpi_items'),[]),normalize_json(r.get('breakdown'),[]),text(r.get('computed_at'))))
        for r in d['Channels']:
            con.execute('INSERT INTO channels(email,slot,platform,page_id,page_name,note) VALUES(?,?,?,?,?,?)',
                (text(r.get('email')).lower(),int(val(r.get('slot')) or 1),text(r.get('platform')),text(r.get('page_id')),text(r.get('page_name')),text(r.get('note'))))
        for r in d['ChannelStats']:
            con.execute('INSERT INTO channel_stats(month,email,slot,videos,views,followers,synced_at,source) VALUES(?,?,?,?,?,?,?,?)',
                (month(r.get('month')),text(r.get('email')).lower(),int(val(r.get('slot')) or 1),val(r.get('videos')) or 0,val(r.get('views')) or 0,val(r.get('followers')) or 0,text(r.get('synced_at')),text(r.get('source'))))
        mapping={'users':counts['Users']+1,'tasks':counts['Tasks'],'content_plan':counts['ContentPlan'],'schedules':counts['Shoots']+counts['Lives'],'meetings':counts['Meetings'],'shifts':counts['Shifts'],'policies':counts['Policy'],'inputs':counts['Inputs'],'payroll':counts['Payroll'],'channels':counts['Channels'],'channel_stats':counts['ChannelStats']}
        actual={table:con.execute(f'SELECT count(*) FROM {table}').fetchone()[0] for table in mapping}
        if actual != mapping: raise ValueError(f'post-import reconciliation mismatch: {actual}')
        con.commit()
        print('Imported and reconciled rows:',json.dumps(actual,sort_keys=True))
        print('Backup:',str(backup))
    except Exception:
        con.rollback()
        raise
    finally: con.close()

def main():
    p=argparse.ArgumentParser()
    sub=p.add_subparsers(dest='command',required=True)
    a=sub.add_parser('prepare'); a.add_argument('xlsx'); a.add_argument('json')
    b=sub.add_parser('apply'); b.add_argument('json'); b.add_argument('--db',required=True)
    args=p.parse_args()
    if args.command=='prepare': prepare(args.xlsx,args.json)
    else: apply(args.json,args.db)

if __name__=='__main__': main()
